# -*- coding: utf-8 -*-
"""
gen_manual_docx.py —— 生成《因子手册与操作指南》Word 版（.docx）
================================================================================
内容与看板（src/factor_docs.py 因子手册 + src/guide.py 框架/指南）保持一致，
从 factor_docs.as_concl(IND) 动态取数，避免文档与代码不一致。
"""
from __future__ import annotations

import os
import re
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

import config as C
import data_loader as dl
import factor_engine as fe
import factor_docs


# ---------------------------------------------------------------------------
# HTML 清理：把 <code>/<b>/<br> 等转成纯文本
# ---------------------------------------------------------------------------
def _clean(s: str) -> str:
    s = re.sub(r"<br\s*/?>", "\n", s)
    s = re.sub(r"</?(code|b|i|strong|em|u)>", "", s)
    s = s.replace("&gt;", ">").replace("&lt;", "<").replace("&amp;", "&")
    s = re.sub(r"\s+", " ", s).strip()
    return s


def _set_font(run, name="微软雅黑", size=10.5, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _heading(doc, text, level=1):
    p = doc.add_paragraph()
    r = p.add_run(text)
    sizes = {0: 20, 1: 15, 2: 12.5, 3: 11}
    colors = {0: "1F2D3D", 1: "1F6FEB", 2: "1F2D3D", 3: "37485C"}
    _set_font(r, "微软雅黑", sizes.get(level, 11), bold=True,
              color=colors.get(level, "1F2D3D"))
    p.paragraph_format.space_before = Pt(10 if level <= 2 else 6)
    p.paragraph_format.space_after = Pt(4)
    return p


def _para(doc, text, size=10.5, bold=False, color=None, indent=0):
    p = doc.add_paragraph()
    r = p.add_run(text)
    _set_font(r, "宋体", size, bold, color)
    p.paragraph_format.left_indent = Cm(indent)
    p.paragraph_format.space_after = Pt(3)
    return p


def _cell_text(cell, text, bold=False, size=10, color=None, bg=None):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(text)
    _set_font(r, "宋体", size, bold, color)
    if bg:
        tcPr = cell._tc.get_or_add_tcPr()
        shd = tcPr.makeelement(qn("w:shd"), {qn("w:val"): "clear",
                                             qn("w:color"): "auto",
                                             qn("w:fill"): bg})
        tcPr.append(shd)


def build():
    b = dl.load_all()
    F = fe.build_all(b)
    IND = F["ind"]
    M = fe.market_summary(IND, F["mkt"], b)
    cards = factor_docs.as_concl(IND)

    doc = Document()
    # 页面边距
    for s in doc.sections:
        s.top_margin = Cm(2.0); s.bottom_margin = Cm(2.0)
        s.left_margin = Cm(2.2); s.right_margin = Cm(2.2)

    # 标题
    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(t.add_run("公募基金拥挤度与集中度模型"), "微软雅黑", 22, True, "1F2D3D")
    t2 = doc.add_paragraph(); t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(t2.add_run("因子手册与操作指南"), "微软雅黑", 16, True, "1F6FEB")
    t3 = doc.add_paragraph(); t3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(t3.add_run("版本 R9（2026-10-03）· 数据源 Wind（主力）+ 东方财富妙想/同花顺 iFinD（交叉验证）"),
              "宋体", 9, False, "6B7A8C")

    # 0 定位
    _heading(doc, "0. 一句话定位", 1)
    _para(doc, "本模型度量「机构资金在行业间拥挤到了什么程度」，是一个风险描述与组合约束工具，"
               "不是择时或收益预测工具。核心结论：拥挤度因子在 36 期长历史下无稳定的截面收益预测力，"
               "正确用法是「风险排序 + 压力测试 + 集中度约束」。", 10.5)

    # 1 因子手册
    _heading(doc, "1. 因子手册（为什么选 · 怎么算 · 衡量什么 · 怎么判读）", 1)
    for c in cards:
        _heading(doc, c.get("title", ""), 2)
        _para(doc, _clean(c.get("concl", "")), 10, False, "37485C")
        why = c.get("why", {})
        if why:
            tb = doc.add_table(rows=len(why), cols=2)
            tb.style = "Table Grid"
            tb.alignment = WD_TABLE_ALIGNMENT.CENTER
            for i, (k, v) in enumerate(why.items()):
                _cell_text(tb.rows[i].cells[0], k, bold=True, size=9.5,
                           bg="EEF3FB")
                _cell_text(tb.rows[i].cells[1], _clean(str(v)), size=9.5)
            for row in tb.rows:
                row.cells[0].width = Cm(3.2)
                row.cells[1].width = Cm(12.0)

    # 2 投资机会判断框架
    _heading(doc, "2. 投资机会判断框架（三层法）", 1)
    _para(doc, "前提声明：36 期 IC 检验显示拥挤度无稳定截面收益预测力，因此本框架不回答「会涨会跌」，"
               "只回答「哪里脆弱、哪里相对安全、哪里需要约束」。", 10.5, bold=True)
    _heading(doc, "第一层 · 风险回避（减配 / 回避）", 3)
    for s in ["极度拥挤：历史分位≥90（F3）或配置倍数>2.5（F12）→ 回避，存量设止损",
              "三信号共振：F10=3/3（持仓+交易+筹码同时极端）→ 强回避",
              "新增拥挤：F15e 高（绝对超配高、但相对历史尚不极端）→ 警惕快速反转"]:
        _para(doc, "• " + s, 10, indent=0.5)
    _heading(doc, "第二层 · 机会识别（关注 / 试探）", 3)
    for s in ["低拥挤：显著低配（分位≤10）→ 列入关注名单",
              "拐点确认：动量 Z（F4b）由负转正 → 资金边际转向",
              "快层佐证：ETF 资金净流入（F14）→ 日频增量确认",
              "满足 ≥2/3（低拥挤+动量转正+资金流入）→ 升级「试探性机会」"]:
        _para(doc, "• " + s, 10, indent=0.5)
    _heading(doc, "第三层 · 组合约束（分散化 / 压力测试）", 3)
    for s in ["集中度红线：HHI>0.15 或 CR3>60% → 主动再平衡",
              "单点压测：单一行业占比≥40% → 强制回撤 30% 拖累测算",
              "K 型分化：极度拥挤≥3 且显著低配≥10 → 回避极端、向低拥挤分散"]:
        _para(doc, "• " + s, 10, indent=0.5)
    _para(doc, "落地决策：先「风险回避」→ 再「机会识别」小仓位试探 → 最后「组合约束」控制敞口。"
               "任何一层都不输出「满仓/清仓」，只输出风险预算的加减。", 10.5, bold=True)

    # 3 操作指南
    _heading(doc, "3. 操作指南（9 页看板怎么看）", 1)
    _heading(doc, "读结论（先看）", 3)
    _para(doc, "① 总览：综合判断（市场状态/HHI分位/极度拥挤数/FCI + 5 条判断理由）+ 投资机会解读", 10, indent=0.5)
    _para(doc, "⑥ 轮动裁决：压力测试结论（可直接用于风控）", 10, indent=0.5)
    _heading(doc, "读数据（深挖）", 3)
    _para(doc, "② 拥挤度总览：单行业钻取 + 31 行业等级明细 + 信号触发记录", 10, indent=0.5)
    _para(doc, "③ 占比变迁：堆叠面积图 / 热力图（标注数值）/ 瀑布图，看资金搬家路径", 10, indent=0.5)
    _para(doc, "④ ETF 资金流：宽基 + 行业 ETF 日度资金动向（快层）", 10, indent=0.5)
    _heading(doc, "读方法（审计）", 3)
    _para(doc, "⑤ 因子研究：因子手册 + 投资机会框架 + 全部检验表；⑦ 稳健性检验；⑧ 数据质量与溯源；⑨ 方法说明", 10, indent=0.5)
    _heading(doc, "读图三原则", 3)
    for s in ["先看「拐点」再看「水平」——拐点（环比转向）比绝对高低更有信息",
              "占比类指标（HHI/CR/超配）跨期可比，绝对规模只在同口径区间内比较",
              "单期极值可能由口径或单期驱动，结论要回 ⑤⑦⑧ 页交叉确认"]:
        _para(doc, "• " + s, 10, indent=0.5)

    # 4 数据更新指南
    _heading(doc, "4. 数据持续更新指南（每季报期怎么做）", 1)
    _para(doc, "更新频率：基金季报每季末后约 15 个工作日披露完毕（4/7/10/次年1 月中下旬），每季更新一次。", 10.5, bold=True)
    _heading(doc, "第 1 步 · 抓取新季报（Wind）", 3)
    for s in ["持仓：主动权益三类（普通股票型/偏股混合型/灵活配置型）前十大重仓股，逐类型查询后相加，必须与「类型总额」交叉验证（gap 约 5~11%）",
              "基准权重：沪深300 成份股按申万一级行业权重，逐期校验加总=100%",
              "行业收益/成交额：申万一级 31 行业季度涨跌幅与成交额，校验 31 行业齐全"]:
        _para(doc, "• " + s, 10, indent=0.5)
    _heading(doc, "第 2 步 · 重建下游（命令）", 3)
    for s in ["python tools/wind_fetch.py build     # 重建超配比例、HHI/CR、因子面板",
              "python src/update_pipeline.py          # 重跑 IC/分层/稳健性，重建 ic_report.xlsx",
              "python tools/seed_daily_raw.py         # 刷新 ETF 日度资金流",
              "python tools/data_quality_report.py    # 刷新数据质量报告"]:
        _para(doc, s, 9.5, indent=0.5, color="37485C")
    _heading(doc, "第 3 步 · 校验验收（必做）", 3)
    for s in ["python tools/apptest_all_pages.py     # 9 页 0 异常",
              "preflight 硬约束 H1–H8 全通过、S1 告警 = 0",
              "检查口径断点：占比类指标跨期可比，绝对规模只在同口径区间比较"]:
        _para(doc, "• " + s, 10, indent=0.5)
    _heading(doc, "数据治理红线", 3)
    for s in ["任何一期数据都要过「三重校验」——31 行业集合完整、总量环比平滑、HHI 落位合理，否则不摄入",
              "Wind 自然语言接口对年报期（Q4）与早年 Q1/Q3 的持仓问句不稳定，这类响应一律不采信并登记根因",
              "数据源优先顺序：Wind（主力）→ 东方财富妙想 / 同花顺 iFinD（交叉验证与补缺）"]:
        _para(doc, "• " + s, 10, indent=0.5)

    # 5 边界声明
    _heading(doc, "5. 已知边界（如实声明）", 1)
    for s in ["前十大重仓口径 ≠ 全部持仓：系统性低估分散型行业，绝对水平有偏（横向可比）",
              "F13 筹码盈利比例为指数层面代理，非成分股加权精确口径",
              "可建模区间主口径仅 10 期（2024Q1~2026Q2）：主口径版 F2/F3/F4b/F15 无法取值，已用偏股混合 36 期超配长历史构建长历史版完成检验（结论：不显著）",
              "早年持仓（2010~2015、2016~2018 Q1/Q4）缺口：需批量基金季报重仓数据库（东方财富妙想为「查实体」接口、套餐对聚合数据已超限；同花顺 iFinD 需另行提供密钥）",
              "结论是「方向性证据」而非「稳健统计结论」，不构成可交易信号"]:
        _para(doc, "• " + s, 10, indent=0.5)

    out = os.path.join(os.path.dirname(__file__), "..", "因子手册与操作指南.docx")
    doc.save(out)
    print("已生成:", os.path.abspath(out))
    return out


if __name__ == "__main__":
    build()
